"""Extracts plain text from uploaded/loaded requirements documents.

Handles plain text/Markdown natively, and optionally extracts text from
``.docx`` and ``.pdf`` files (common formats for PRDs/BRDs) via the optional
``docs`` extra (``pip install -e ".[docs]"``). Anything else that doesn't
decode into readable text (e.g. a binary/proprietary export accidentally
uploaded) raises :class:`UnsupportedDocumentError` with a clear message,
instead of silently feeding garbage bytes into the parser.
"""

from __future__ import annotations

import io
from pathlib import Path

_PLAIN_TEXT_EXTENSIONS = {".md", ".markdown", ".txt", ".text", ""}


class UnsupportedDocumentError(RuntimeError):
    """Raised when a requirements document can't be read as text."""


def extract_text(filename: str, data: bytes) -> str:
    """Return the textual content of *data* (the bytes of *filename*).

    Raises :class:`UnsupportedDocumentError` if the format is not supported
    or the extractor dependency is missing.
    """

    suffix = Path(filename).suffix.lower()

    if suffix == ".docx":
        return _extract_docx(data, filename)
    if suffix == ".pdf":
        return _extract_pdf(data, filename)

    text = data.decode("utf-8", errors="replace")
    if suffix not in _PLAIN_TEXT_EXTENSIONS:
        # Unrecognized extension: only accept it if it actually looks like text.
        _require_looks_like_text(text, filename)
    else:
        _require_looks_like_text(text, filename)
    return text


def _require_looks_like_text(text: str, filename: str) -> None:
    sample = text[:4000]
    if not sample.strip():
        return
    printable = sum(1 for c in sample if c.isprintable() or c in "\n\r\t")
    replacement_chars = sample.count("\ufffd")
    if printable / len(sample) < 0.85 or replacement_chars > len(sample) * 0.05:
        raise UnsupportedDocumentError(
            f"'{filename}' doesn't look like a plain-text or Markdown document "
            "(it may be a binary/proprietary format such as .docx, .pdf, or an "
            "exported Confluence/Word file that wasn't saved as plain text). "
            "Please paste the requirements text directly, save it as .md/.txt, "
            "or upload a .docx/.pdf file (requires `pip install -e \".[docs]\"`)."
        )


def _extract_docx(data: bytes, filename: str) -> str:
    try:
        import docx
    except ImportError as exc:
        raise UnsupportedDocumentError(
            f"Reading '{filename}' requires the 'python-docx' package. "
            'Install it with: pip install -e ".[docs]"'
        ) from exc

    try:
        document = docx.Document(io.BytesIO(data))
    except Exception as exc:
        raise UnsupportedDocumentError(f"Could not read '{filename}' as a .docx file: {exc}") from exc

    parts: list[str] = []
    for para in document.paragraphs:
        text = para.text.strip()
        if not text:
            continue
        style_name = (para.style.name if para.style else "") or ""
        if style_name.lower().startswith("heading") or style_name.lower() == "title":
            digits = "".join(ch for ch in style_name if ch.isdigit())
            level = max(1, min(int(digits), 3)) if digits else 1
            parts.append(f"{'#' * level} {text}")
        else:
            parts.append(text)

    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            if any(cells):
                parts.append("- " + " | ".join(cells))

    text = "\n\n".join(parts)
    if not text.strip():
        raise UnsupportedDocumentError(f"No readable text could be extracted from '{filename}'.")
    return text


def _extract_pdf(data: bytes, filename: str) -> str:
    try:
        from pypdf import PdfReader
    except ImportError as exc:
        raise UnsupportedDocumentError(
            f"Reading '{filename}' requires the 'pypdf' package. "
            'Install it with: pip install -e ".[docs]"'
        ) from exc

    try:
        reader = PdfReader(io.BytesIO(data))
        pages_text = [page.extract_text() or "" for page in reader.pages]
    except Exception as exc:
        raise UnsupportedDocumentError(f"Could not read '{filename}' as a .pdf file: {exc}") from exc

    text = "\n\n".join(pages_text)
    if not text.strip():
        raise UnsupportedDocumentError(
            f"No readable text could be extracted from '{filename}' (it may be a scanned/image-only PDF)."
        )
    return text
