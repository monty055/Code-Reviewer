"""Renders a :class:`~reviewer_agent.test_case_models.TestCaseSuite` into the
formats a QA team typically wants to consume: Markdown, JSON, CSV, a Word
(``.docx``) table, or a PDF table.

CSV/Markdown/JSON have no extra dependencies. DOCX export reuses
``python-docx`` (the ``docs`` extra, already used to *read* .docx
requirements documents). PDF export uses ``reportlab`` (the ``export``
extra) and raises a clear, actionable error -- consistent with how missing
optional dependencies are handled elsewhere in this project -- if it isn't
installed.
"""

from __future__ import annotations

import csv
import io
import json as _json
from datetime import datetime, timezone

from .test_case_models import TestCase, TestCaseSuite

_COLUMNS = [
    "ID",
    "Feature",
    "Criterion",
    "Type",
    "Priority",
    "Title",
    "Preconditions",
    "Steps",
    "Test Data",
    "Expected Result",
    "Status",
    "Open Questions",
]


class ExportDependencyError(RuntimeError):
    """Raised when an export format's optional dependency isn't installed."""


def _row_values(tc: TestCase) -> list[str]:
    test_data = "; ".join(f"{row.field} = {row.value}" for row in tc.test_data)
    return [
        tc.id,
        f"{tc.feature_id}: {tc.feature_title}",
        f"{tc.criterion_id}: {tc.criterion_text}",
        tc.type.value,
        tc.priority.value,
        tc.title,
        " | ".join(tc.preconditions),
        " -> ".join(tc.steps),
        test_data,
        tc.expected_result,
        "Ready" if tc.is_ready else "Needs Clarification",
        " | ".join(tc.open_questions),
    ]


def render_markdown(suite: TestCaseSuite) -> str:
    generated_at = datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M UTC")
    lines: list[str] = []
    lines.append("# Generated Test Cases")
    lines.append("")
    lines.append(f"- **Requirements document:** `{suite.requirements_source}`")
    lines.append(f"- **Generated:** {generated_at}")
    lines.append(f"- **Total test cases:** {len(suite.test_cases)}")
    lines.append(
        f"- **Ready to automate/execute as-is:** {suite.ready_count}/{len(suite.test_cases)} "
        f"({suite.readiness_score:.0%})"
    )
    lines.append("")

    for tc in suite.test_cases:
        status = "✅ Ready" if tc.is_ready else "🟠 Needs Clarification"
        lines.append(f"## {tc.id}: {tc.title}")
        lines.append("")
        lines.append(f"- **Feature:** {tc.feature_id} — {tc.feature_title}")
        lines.append(f"- **Criterion:** {tc.criterion_id} — {tc.criterion_text}")
        lines.append(f"- **Type:** {tc.type.value} · **Priority:** {tc.priority.value} · **Status:** {status}")
        lines.append(f"- **Preconditions:** {' | '.join(tc.preconditions)}")
        lines.append("- **Steps:**")
        for i, step in enumerate(tc.steps, start=1):
            lines.append(f"  {i}. {step}")
        lines.append("- **Test data:**")
        for row in tc.test_data:
            lines.append(f"  - `{row.field}` = `{row.value}`")
        lines.append(f"- **Expected result:** {tc.expected_result}")
        if tc.open_questions:
            lines.append("- **Open questions (blocking automation):**")
            for q in tc.open_questions:
                lines.append(f"  - {q}")
        lines.append("")

    return "\n".join(lines) + "\n"


def render_json(suite: TestCaseSuite) -> str:
    payload = {
        "requirements_source": suite.requirements_source,
        "total_test_cases": len(suite.test_cases),
        "ready_count": suite.ready_count,
        "readiness_score": suite.readiness_score,
        "test_cases": [
            {
                "id": tc.id,
                "feature_id": tc.feature_id,
                "feature_title": tc.feature_title,
                "criterion_id": tc.criterion_id,
                "criterion_text": tc.criterion_text,
                "title": tc.title,
                "type": tc.type.value,
                "priority": tc.priority.value,
                "preconditions": tc.preconditions,
                "steps": tc.steps,
                "test_data": [{"field": r.field, "value": r.value} for r in tc.test_data],
                "expected_result": tc.expected_result,
                "is_ready": tc.is_ready,
                "open_questions": tc.open_questions,
            }
            for tc in suite.test_cases
        ],
    }
    return _json.dumps(payload, indent=2)


def render_csv(suite: TestCaseSuite) -> str:
    buf = io.StringIO()
    writer = csv.writer(buf)
    writer.writerow(_COLUMNS)
    for tc in suite.test_cases:
        writer.writerow(_row_values(tc))
    return buf.getvalue()


def render_docx(suite: TestCaseSuite) -> bytes:
    try:
        import docx
    except ImportError as exc:
        raise ExportDependencyError(
            "Exporting to .docx requires the 'python-docx' package. Install it with: pip install -e \".[docs]\""
        ) from exc

    document = docx.Document()
    document.add_heading("Generated Test Cases", level=1)
    document.add_paragraph(f"Requirements document: {suite.requirements_source}")
    document.add_paragraph(
        f"Total test cases: {len(suite.test_cases)} · Ready: {suite.ready_count} "
        f"({suite.readiness_score:.0%})"
    )

    table = document.add_table(rows=1, cols=len(_COLUMNS))
    table.style = "Light Grid Accent 1"
    header_cells = table.rows[0].cells
    for i, col in enumerate(_COLUMNS):
        header_cells[i].text = col

    for tc in suite.test_cases:
        cells = table.add_row().cells
        for i, value in enumerate(_row_values(tc)):
            cells[i].text = value

    buf = io.BytesIO()
    document.save(buf)
    return buf.getvalue()


def render_pdf(suite: TestCaseSuite) -> bytes:
    try:
        from reportlab.lib import colors
        from reportlab.lib.pagesizes import landscape, letter
        from reportlab.lib.styles import getSampleStyleSheet
        from reportlab.platypus import Paragraph, SimpleDocTemplate, Table, TableStyle
    except ImportError as exc:
        raise ExportDependencyError(
            "Exporting to .pdf requires the 'reportlab' package. Install it with: pip install -e \".[export]\""
        ) from exc

    styles = getSampleStyleSheet()
    cell_style = styles["BodyText"]
    cell_style.fontSize = 7
    header_style = styles["BodyText"].clone("header")
    header_style.fontSize = 8
    header_style.textColor = colors.white

    def cell(text: str, header: bool = False) -> Paragraph:
        style = header_style if header else cell_style
        return Paragraph(text.replace("\n", "<br/>"), style)

    data = [[cell(col, header=True) for col in _COLUMNS]]
    for tc in suite.test_cases:
        data.append([cell(v) for v in _row_values(tc)])

    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=landscape(letter), title="Generated Test Cases")
    table = Table(data, repeatRows=1)
    table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#3355dd")),
                ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f2f4fb")]),
            ]
        )
    )

    title_style = styles["Heading1"]
    subtitle_style = styles["Normal"]
    elements = [
        Paragraph("Generated Test Cases", title_style),
        Paragraph(f"Requirements document: {suite.requirements_source}", subtitle_style),
        Paragraph(
            f"Total test cases: {len(suite.test_cases)} · Ready: {suite.ready_count} "
            f"({suite.readiness_score:.0%})",
            subtitle_style,
        ),
        table,
    ]
    doc.build(elements)
    return buf.getvalue()
